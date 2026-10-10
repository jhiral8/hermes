import datetime
import io
import json
import os
import tempfile
import unittest
import urllib.error
from pathlib import Path
from types import SimpleNamespace

import openpyxl

import import_macrofactor as imp


def export(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Calories & Macros"
    ws.append(["Date", "Calories (kcal)", "Fat (g)", "Carbs (g)", "Protein (g)"])
    ws.append([datetime.datetime(2026, 8, 19), 2037, 62, 192, 185])
    ws.append([datetime.datetime(2026, 8, 20), 2496, 68, 226, 245])
    ws.append([datetime.datetime(2026, 8, 21), None, None, None, None])  # no total: skipped
    ws.append([datetime.datetime(2026, 8, 22), 2100, 60, 200, 150])      # last day: skipped
    p = wb.create_sheet("Partial Logging")
    p.append(["Date", "Partial"])
    p.append(["20/08/2026", "Yes"])
    p.append(["19/08/2026", "No"])
    w = wb.create_sheet("Scale Weight")
    w.append(["Date", "Weight (kg)", "Fat Percent"])
    w.append([datetime.datetime(2026, 8, 19), 86.0, None])
    w.append([datetime.datetime(2026, 8, 20), 84.1, None])
    w.append([datetime.datetime(2026, 8, 22), None, None])
    wb.save(path)
    return path


class Resp(io.BytesIO):
    status = 200


class Fake:
    def __init__(self, fail=None):
        self.calls, self.fail = [], fail

    def __call__(self, req, timeout=None):
        self.calls.append((req.get_method(), req.full_url, req.get_header("Authorization"), json.loads(req.data)))
        if self.fail:
            raise urllib.error.HTTPError(req.full_url, self.fail, "no", {}, io.BytesIO(b""))
        return Resp(b"{}")


def args(path, **kw):
    d = tempfile.mkdtemp()
    base = dict(export=path, base="http://nt.test", key_file=None, state=os.path.join(d, "state.json"),
                backup_dir=d, pause=0, apply=False)
    base.update(kw)
    return SimpleNamespace(**base), d


def backup(d):
    Path(d, "nutritrace-20261008.db").write_bytes(b"SQLite format 3\x00")


class Plan(unittest.TestCase):
    def setUp(self):
        self.path = export(os.path.join(tempfile.mkdtemp(), "mf.xlsx"))

    def test_last_day_and_empty_days_are_left_out(self):
        totals, weights, skipped, partial = imp.read_export(self.path)
        self.assertEqual(sorted(totals), [datetime.date(2026, 8, 19)])  # 20 Aug partly logged, 22 Aug last day
        self.assertEqual(skipped, datetime.date(2026, 8, 22))
        self.assertEqual(partial, [datetime.date(2026, 8, 20)])
        self.assertEqual(weights, {datetime.date(2026, 8, 19): 86.0, datetime.date(2026, 8, 20): 84.1})

    def test_plan_has_stable_keys_and_the_expected_bodies(self):
        totals, weights, _, _ = imp.read_export(self.path)
        a = imp.plan(totals, weights)
        b = imp.plan(totals, weights)
        self.assertEqual([i["key"] for i in a], [i["key"] for i in b])
        self.assertEqual(len({i["key"] for i in a}), len(a))  # no duplicate keys
        first = a[0]
        self.assertEqual(first["body"]["nutrition"], {"calories": 2037, "fat": 62, "carbohydrates": 192, "proteins": 185})
        self.assertEqual(len(a), 3)  # one food entry and two weigh-ins
        self.assertEqual(first["body"]["name"], imp.FOOD_NAME)


class Run(unittest.TestCase):
    def setUp(self):
        self.path = export(os.path.join(tempfile.mkdtemp(), "mf.xlsx"))

    def test_dry_run_writes_nothing(self):
        a, d = args(self.path)
        fake = Fake()
        out = io.StringIO()
        imp.run(a, client=imp.Client(a.base, "k", opener=fake, pause=0), out=out)
        self.assertEqual(fake.calls, [])
        self.assertFalse(os.path.exists(a.state))
        self.assertIn("Dry run", out.getvalue())
        self.assertIn("1 food entries, 2 weigh-ins", out.getvalue())

    def test_apply_refuses_without_a_backup(self):
        a, d = args(self.path, apply=True)
        with self.assertRaises(SystemExit) as e:
            imp.run(a, client=imp.Client(a.base, "k", opener=Fake(), pause=0), out=io.StringIO())
        self.assertIn("Refusing to apply", str(e.exception))

    def test_apply_writes_with_the_key_and_the_right_routes(self):
        a, d = args(self.path, apply=True)
        backup(d)
        fake = Fake()
        imp.run(a, client=imp.Client(a.base, "nt_write", opener=fake, pause=0), out=io.StringIO())
        methods = [(m, u.split("/api/v1")[1]) for m, u, _, _ in fake.calls]
        self.assertEqual(methods, [("POST", "/diary/2026-08-19/food"),
                                   ("PUT", "/diary/2026-08-19/body-stat"), ("PUT", "/diary/2026-08-20/body-stat")])
        self.assertTrue(all(auth == "Bearer nt_write" for _, _, auth, _ in fake.calls))
        self.assertEqual(fake.calls[1][3], {"weight": 86.0})

    def test_rerun_skips_what_was_already_written(self):
        a, d = args(self.path, apply=True)
        backup(d)
        fake = Fake()
        imp.run(a, client=imp.Client(a.base, "k", opener=fake, pause=0), out=io.StringIO())
        again = Fake()
        out = io.StringIO()
        imp.run(a, client=imp.Client(a.base, "k", opener=again, pause=0), out=out)
        self.assertEqual(again.calls, [])
        self.assertIn("Already loaded: 3", out.getvalue())

    def test_a_failure_stops_the_run_and_keeps_what_was_written(self):
        a, d = args(self.path, apply=True)
        backup(d)
        calls = []

        def flaky(req, timeout=None):
            calls.append(req)
            if len(calls) == 2:
                raise urllib.error.HTTPError(req.full_url, 403, "no", {}, io.BytesIO(b""))
            return Resp(b"{}")
        with self.assertRaises(RuntimeError) as e:
            imp.run(a, client=imp.Client(a.base, "k", opener=flaky, pause=0), out=io.StringIO())
        self.assertIn("403", str(e.exception))
        self.assertEqual(len(imp.load_state(a.state)), 1)  # the write before the failure is recorded
        self.assertNotIn("Bearer", str(e.exception))  # no key in the message

    def test_pacing_waits_between_writes(self):
        waits = []
        c = imp.Client("http://nt.test", "k", opener=Fake(), pause=1.1, sleep=waits.append)
        item = {"kind": "weight", "date": "2026-08-19", "body": {"weight": 86.0}}
        c.write(item)
        c.write(item)
        self.assertEqual(len(waits), 1)
        self.assertGreater(waits[0], 0)


if __name__ == "__main__":
    unittest.main()
