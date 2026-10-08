import datetime
import json
import tempfile
import unittest
from pathlib import Path

import openpyxl

import history_feed as hf
from sources import SourceError


def make_export(days, weights=None, partial=()):
    """A small MacroFactor-shaped .xlsx: days is [(date, kcal)], weights {date: kg}."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Calories & Macros"
    ws.append(["Date", "Calories (kcal)", "Fat (g)", "Carbs (g)", "Protein (g)"])
    for d, kcal in days:
        ws.append([datetime.datetime.combine(d, datetime.time()), kcal, 1, 1, 1])
    p = wb.create_sheet("Partial Logging")
    p.append(["Date", "Partial"])
    for d in partial:
        p.append([d.strftime("%d/%m/%Y"), "Yes"])
    w = wb.create_sheet("Scale Weight")
    w.append(["Date", "Weight (kg)", "Fat Percent"])
    for d, kg in (weights or {}).items():
        w.append([datetime.datetime.combine(d, datetime.time()), kg, None])
    path = Path(tempfile.mkdtemp(), "MacroFactor.xlsx")
    wb.save(path)
    return path

D = datetime.date


def history(days, weights=None):
    return ({d["date"]: {"kcal": d["kcal"], "status": d["status"]} for d in days}, weights or {})


class Export(unittest.TestCase):
    def test_partial_and_last_days_are_marked(self):
        p = make_export([(D(2026, 10, 1), 2000), (D(2026, 10, 2), 1800), (D(2026, 10, 3), 700)],
                        {D(2026, 10, 1): 80.0}, partial=[D(2026, 10, 2)])
        days, weights = hf.read_export(p)
        self.assertEqual(days["2026-10-01"], {"kcal": 2000.0, "status": "complete"})
        self.assertEqual(days["2026-10-02"]["status"], "partial")
        self.assertEqual(days["2026-10-03"]["status"], "partial")  # the export's own last day
        self.assertEqual(weights, {"2026-10-01": 80.0})


class Build(unittest.TestCase):
    def test_nutritrace_wins_on_overlap_and_mf_fills_the_rest(self):
        h = history([{"date": "2026-10-01", "kcal": 2000, "status": "complete"},
                     {"date": "2026-10-02", "kcal": 1800, "status": "complete"}], {"2026-10-01": 80.0})
        nt = {"2026-10-02": {"kcal": 2100, "items": 3}}
        out = {r["date"]: r for r in hf.build_days(h, nt, D(2026, 10, 3))}
        self.assertEqual(out["2026-10-01"]["intake"], 2000)
        self.assertEqual(out["2026-10-01"]["weight"], 80.0)
        self.assertEqual(out["2026-10-02"]["intake"], 2100)  # NutriTrace wins
        self.assertEqual(out["2026-10-02"]["status"], "complete")

    def test_nutritrace_day_with_no_items_is_not_a_logged_day(self):
        h = history([{"date": "2026-10-01", "kcal": 2000, "status": "complete"}])
        nt = {"2026-10-01": {"kcal": 0, "items": 0}}
        out = hf.build_days(h, nt, D(2026, 10, 2))
        self.assertEqual(out[0]["intake"], 2000)

    def test_days_with_no_data_are_missing_and_today_is_partial(self):
        h = history([{"date": "2026-10-01", "kcal": 2000, "status": "complete"}])
        out = hf.build_days(h, {"2026-10-03": {"kcal": 900, "items": 2}}, D(2026, 10, 3))
        self.assertEqual([r["status"] for r in out], ["complete", "missing", "partial"])
        self.assertIsNone(out[1]["intake"])
        self.assertEqual(out[2]["intake"], 900)

    def test_mf_partial_days_stay_partial(self):
        h = history([{"date": "2026-10-01", "kcal": 700, "status": "partial"}])
        out = hf.build_days(h, {}, D(2026, 10, 2))
        self.assertEqual(out[0]["status"], "partial")

    def test_no_history_gives_no_days(self):
        self.assertEqual(hf.build_days(({}, {}), {}, D(2026, 10, 2)), [])


class Fetch(unittest.TestCase):
    def test_reads_each_day_and_skips_unreadable_ones(self):
        seen = []

        class FakeApp:
            def __init__(self, *a, **k):
                pass

            def get(self, path, ttl=0):
                seen.append(path)
                if "10-02" in path:
                    raise SourceError("NutriTrace refused the token")
                return {"totals": {"calories": 1500}, "item_count": 2}
        orig = hf.TraceApp
        hf.TraceApp = FakeApp
        try:
            pauses = []
            out = hf.fetch_nutritrace("http://nt.test", "/k", D(2026, 10, 1), D(2026, 10, 3),
                                      pause=1.1, sleep=pauses.append)
        finally:
            hf.TraceApp = orig
        self.assertEqual(sorted(out), ["2026-10-01", "2026-10-03"])
        self.assertEqual(out["2026-10-01"], {"kcal": 1500, "items": 2})
        self.assertEqual(len(pauses), 2)  # paced between calls, not before the first


class Run(unittest.TestCase):
    def test_end_to_end_writes_the_feed_and_the_estimate(self):
        d = Path(tempfile.mkdtemp())
        base = D(2026, 9, 1)
        days = [(base + datetime.timedelta(days=i), 2200) for i in range(30)]
        weights = {base + datetime.timedelta(days=i): 85.0 - 0.05 * i for i in range(0, 30, 3)}
        export = make_export(days, weights)
        res = hf.run(str(export), None, None, str(d / "days.json"), str(d / "estimator.json"),
                     today=D(2026, 9, 30), fetch=lambda *a: {})
        self.assertEqual(res["days"], 30)
        feed = json.loads((d / "days.json").read_text())
        self.assertEqual(feed[0]["weight"], 85.0)
        est = json.loads((d / "estimator.json").read_text())
        self.assertIn("held", est)
        self.assertIn("as_of", est)

    def test_empty_history_stops(self):
        d = Path(tempfile.mkdtemp())
        empty = make_export([])
        with self.assertRaises(SystemExit):
            hf.run(str(empty), None, None, str(d / "a"), str(d / "b"), today=D(2026, 10, 2))


def make_nutritrace_db(diary=(), wellness=()):
    """A small database with NutriTrace's table layout (names from its schema dump)."""
    import sqlite3
    path = Path(tempfile.mkdtemp(), "nutritrace.db")
    con = sqlite3.connect(path)
    con.execute("CREATE TABLE diary (id TEXT, user_id TEXT, date TEXT, items TEXT, body_stats TEXT, "
                "deleted_at TEXT)")
    con.execute("CREATE TABLE wellness_data (id TEXT, user_id TEXT, date TEXT, source TEXT, "
                "metric_type TEXT, value REAL, synced_at TEXT)")
    for row in diary:
        con.execute("INSERT INTO diary VALUES (?, 'u', ?, ?, ?, ?)", row)
    for row in wellness:
        con.execute("INSERT INTO wellness_data VALUES (?, 'u', ?, 'scale', ?, ?, ?)", row)
    con.commit()
    con.close()
    return path


class NutriTraceDb(unittest.TestCase):
    def test_reads_totals_items_and_weights(self):
        items = json.dumps([{"name": "Oats", "nutrition": {"calories": 300}},
                            {"name": "Milk", "nutrition": {"calories": 120.5}},
                            {"name": "?"}])
        p = make_nutritrace_db(
            diary=[("a", "2026-10-02", items, None, None),
                   ("b", "2026-10-03", json.dumps([]), json.dumps({"weight": 81.2}), None),
                   ("c", "2026-10-04", "not json", None, None),
                   ("d", "2026-10-05", json.dumps([{"nutrition": {"calories": 900}}]), None, "2026-10-05")],
            wellness=[("w1", "2026-10-03", "weight_kg", 80.9, "2026-10-03T07:00"),
                      ("w2", "2026-10-03", "weight_kg", 80.7, "2026-10-03T08:00"),
                      ("w3", "2026-10-04", "steps", 9000, "2026-10-04T08:00"),
                      ("w4", "2026-10-05", "weight", 70.0, "2026-10-05T08:00")])
        days, weights = hf.read_nutritrace_db(p, D(2026, 10, 1), D(2026, 10, 8))
        self.assertEqual(days["2026-10-02"], {"kcal": 420.5, "items": 3})
        self.assertEqual(days["2026-10-03"], {"kcal": None, "items": 0})
        self.assertEqual(days["2026-10-04"], {"kcal": None, "items": 0})  # unreadable items: no total
        self.assertNotIn("2026-10-05", days)  # deleted
        self.assertEqual(weights["2026-10-03"], 80.7)  # the later sync wins over body stats
        self.assertNotIn("2026-10-04", weights)  # steps are not weight
        self.assertNotIn("2026-10-05", weights)  # only NutriTrace's weight_kg metric counts

    def test_body_stats_weight_is_converted_from_lb(self):
        self.assertAlmostEqual(hf._weight_of(json.dumps({"weight": 180, "weight_unit": "lb"})), 81.646, places=2)
        self.assertEqual(hf._weight_of(json.dumps({"weight": 81.2, "weight_unit": "kg"})), 81.2)
        self.assertEqual(hf._weight_of(json.dumps({"weight": 81.2})), 81.2)  # kg is the default

    def test_opened_read_only(self):
        import sqlite3
        p = make_nutritrace_db()
        hf.read_nutritrace_db(p, D(2026, 10, 1), D(2026, 10, 8))
        con = sqlite3.connect(p.resolve().as_uri() + "?mode=ro", uri=True)
        with self.assertRaises(sqlite3.OperationalError):
            con.execute("INSERT INTO diary VALUES ('x', 'u', '2026-10-01', '[]', NULL, NULL)")
        con.close()

    def test_run_merges_database_weights_and_food_with_nutritrace_winning(self):
        d = Path(tempfile.mkdtemp())
        base = D(2026, 9, 28)
        days = [(base + datetime.timedelta(days=i), 2100) for i in range(6)]  # 28 Sep to 3 Oct
        export = make_export(days, {D(2026, 9, 28): 84.0})
        p = make_nutritrace_db(
            diary=[("a", "2026-10-01", json.dumps([{"nutrition": {"calories": 1500}}]), None, None)],
            wellness=[("w", "2026-10-02", "weight_kg", 83.4, "2026-10-02T07:00")])
        hf.run(str(export), None, None, str(d / "days.json"), str(d / "estimator.json"),
               today=D(2026, 10, 3), fetch=lambda *a: {}, nutritrace_db=str(p))
        feed = {r["date"]: r for r in json.loads((d / "days.json").read_text())}
        self.assertEqual(feed["2026-10-01"]["intake"], 1500)  # NutriTrace wins on food
        self.assertEqual(feed["2026-10-01"]["status"], "complete")
        self.assertEqual(feed["2026-10-02"]["weight"], 83.4)  # NutriTrace's weigh-in is used
        self.assertEqual(feed["2026-10-02"]["intake"], 2100)  # no NutriTrace food: MacroFactor fills in
        self.assertEqual(feed["2026-09-28"]["weight"], 84.0)
        self.assertEqual(feed["2026-10-03"]["status"], "partial")  # today


if __name__ == "__main__":
    unittest.main()
